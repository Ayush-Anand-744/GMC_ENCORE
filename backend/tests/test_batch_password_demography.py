"""Integration checks for real user-provided cases plus true encrypted fixtures."""
from datetime import date
from pathlib import Path
from io import BytesIO
from collections import Counter

import pytest
from fastapi.testclient import TestClient
from backend.main import app
from backend.orchestration.orchestrator import Orchestrator
from backend.core.demography_engine import recalculate_ages, summary
from backend.core.password_handler import is_password_protected, unlock_bytes, IncorrectPassword, UnsupportedEncryption

FIXTURES=Path(__file__).parent/'fixtures'
SECRET='GMC-Test-2026'

def src(name):return (name,(FIXTURES/name).read_bytes())

def test_real_samples_and_multi_roster_aggregation():
    rfq, emp=src('RFQ_Sample.xlsx'),src('Employee_Data_Sample.xlsx')
    original=Orchestrator().process([rfq,emp],date(2026,9,22))
    assert original['classifications'][0]['category']=='rfq'
    assert original['classifications'][1]['category']=='demography'
    assert len(original['demography'])==36
    assert summary(original['demography'])=={'group_size':36,'primary_members':14,'sme':'YES'}
    rows={r['field']:r for r in original['quote_rows']}
    for key, expected in [('Group Size','36'),('Number of Primary Members','14'),('Flagging SME','YES')]:
        assert rows[key]['coverage_details']==expected
    assert original['demography'][3]['Date of Birth']=='Not Available' # source has 2985
    assert original['demography'][3]['EmpCode']=='Not Available'
    # Independently named copies are exact duplicate uploads; only one copy
    # should survive, while genuinely new populations must be appended.
    copies=Orchestrator().process([rfq,emp,('copy.xlsx',emp[1])],date(2026,9,22))
    assert len(copies['demography'])==36
    # Introduce one truly new member in another workbook.
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(emp[1]));sh=wb.active
    sh.cell(row=3,column=2).value='NEW_EMP'
    sh.cell(row=3,column=3).value='New Unique Employee'
    buf=BytesIO();wb.save(buf)
    extra=Orchestrator().process([rfq,emp,('additional.xlsx',buf.getvalue())],date(2026,9,22))
    assert len(extra['demography'])==37
    assert summary(extra['demography'])['primary_members']==15


def test_encrypted_pdf_real_wrong_password_retry_and_extract():
    filename, encrypted=src('RFQ_Sample_Protected.pdf')
    assert is_password_protected(filename,encrypted)
    with pytest.raises(IncorrectPassword):unlock_bytes(filename,encrypted,'incorrect')
    unlocked=unlock_bytes(filename,encrypted,SECRET)
    assert not is_password_protected(filename,unlocked)
    from backend.core.file_reader import read_file
    result=read_file(filename,unlocked)
    assert 'PHANTOM' in result.text
    client=TestClient(app)
    files=[('files',(filename,encrypted,'application/pdf')),
           ('files',('Employee_Data_Sample.xlsx',(FIXTURES/'Employee_Data_Sample.xlsx').read_bytes(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))]
    bad=client.post('/api/process',files=files,data={'passwords_json':'{"RFQ_Sample_Protected.pdf":"incorrect"}'})
    assert bad.status_code==401 and bad.json()['detail']['locked_file']==filename
    good=client.post('/api/process',files=files,data={'passwords_json':'{"RFQ_Sample_Protected.pdf":"GMC-Test-2026"}'})
    assert good.status_code==200,good.text[:300]
    response=client.get('/api/cases/'+good.json()['case_id'])
    assert response.status_code==200 and len(response.json()['demography'])==36


def test_encrypted_excel_real_wrong_and_correct_password():
    filename, encrypted=src('Employee_Data_Sample_Protected.xlsx')
    assert is_password_protected(filename,encrypted)
    try:
        import msoffcrypto
    except ImportError:
        pytest.skip('Office decrypt integration requires msoffcrypto-tool (declared in requirements.txt)')
    with pytest.raises(IncorrectPassword):unlock_bytes(filename,encrypted,'incorrect')
    unlocked=unlock_bytes(filename,encrypted,SECRET)
    assert unlocked[:2]==b'PK'
    from backend.core.file_reader import read_file
    rr=read_file(filename,unlocked)
    assert len(rr.tables)==1
    assert len(Orchestrator().process([('roster.xlsx',unlocked)],date(2026,9,22))['demography'])==36
    client=TestClient(app)
    locked_files=[('files',(filename,encrypted,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'))]
    bad=client.post('/api/process',files=locked_files,data={'passwords_json':'{"Employee_Data_Sample_Protected.xlsx":"incorrect"}'})
    assert bad.status_code==401, bad.text
    good=client.post('/api/process',files=locked_files,data={'passwords_json':'{"Employee_Data_Sample_Protected.xlsx":"GMC-Test-2026"}'})
    assert good.status_code==200,good.text[:300]
    case=client.get('/api/cases/'+good.json()['case_id']).json()
    assert len(case['demography'])==36


def test_date_mode_member_changes_and_download_consistency():
    client=TestClient(app)
    filename, data=src('Employee_Data_Sample.xlsx')
    post=client.post('/api/process', files={'files':(filename,data,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
    assert post.status_code==200,post.text[:300]
    case_id=post.json()['case_id'];url=f'/api/cases/{case_id}'
    initial=client.get(url).json()
    first=initial['demography'][0]
    updated=client.post(url+'/reference-date',data={'use_policy_date':'true','policy_date':'2027-04-04'})
    assert updated.status_code==200
    assert updated.json()['demography'][0]['Age']==35
    updated=client.post(url+'/reference-date',data={'use_policy_date':'true','policy_date':'2027-04-05'})
    assert updated.status_code==200
    assert updated.json()['demography'][0]['Age']==36
    rows=updated.json()['demography'];rows += [dict(rows[0],Name='Another Primary',EmpCode='E_NEW')]
    edit=client.put(url+'/demography',json=rows)
    assert edit.status_code==200,edit.text[:300]
    edited=edit.json();assert len(edited['demography'])==37
    fields={x['field']:x['coverage_details'] for x in edited['quote_rows']}
    assert fields['Group Size']=='37' and fields['Number of Primary Members']=='15' and fields['Flagging SME']=='YES'
    dl=client.get(url+'/download/demography')
    assert dl.status_code==200
    import pandas as pd
    df=pd.read_excel(BytesIO(dl.content))
    assert len(df)==37 and df.iloc[0]['Age']==36
    no_demo=client.put(url+'/demography',json=[])
    assert no_demo.status_code==200
    fields={x['field']:x['coverage_details'] for x in no_demo.json()['quote_rows']}
    assert fields['Group Size']=='Not Available' and fields['Flagging SME']=='Not Available'
