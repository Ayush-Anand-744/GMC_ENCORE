from io import BytesIO
from pathlib import Path
import email
import pandas as pd
from pypdf import PdfReader
from docx import Document
from PIL import Image

class ReadResult:
    def __init__(self, text="", tables=None, metadata=None):
        self.text = text
        self.tables = tables or []
        self.metadata = metadata or {}


def _excel(data: bytes) -> ReadResult:
    book = pd.read_excel(BytesIO(data), sheet_name=None, header=None, dtype=object)
    tables = []
    chunks = []
    for name, df in book.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        tables.append((name, df))
        sheet_text = [f"[SHEET: {name}]"]
        for _, row in df.iterrows():
            vals = [str(v).strip() for v in row if pd.notna(v) and str(v).strip() != ""]
            if not vals: continue
            if len(vals) == 1:
                sheet_text.append(vals[0])
            elif len(vals) == 2:
                sheet_text.append(f"{vals[0]}: {vals[1]}")
            else:
                sheet_text.append(f"{vals[0]}: {vals[1]} | {vals[2]}")
        chunks.append("\n".join(sheet_text))
    return ReadResult("\n\n".join(chunks), tables, {"sheets": list(book)})



def _pdf(data: bytes) -> ReadResult:
    reader = PdfReader(BytesIO(data))
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    return ReadResult(text, metadata={"pages": len(reader.pages), "needs_ocr": len(text.strip()) < 80})


def _docx(data: bytes) -> ReadResult:
    doc = Document(BytesIO(data))
    chunks = [p.text for p in doc.paragraphs]
    tables = []
    for idx, t in enumerate(doc.tables):
        rows = [[c.text for c in r.cells] for r in t.rows]
        if rows:
            df = pd.DataFrame(rows[1:], columns=rows[0] if rows[0] else None)
            tables.append((f"Table {idx+1}", df))
            chunks.extend(" | ".join(r) for r in rows)
    return ReadResult("\n".join(chunks), tables)


def _eml(data: bytes) -> ReadResult:
    msg = email.message_from_bytes(data)
    chunks = [f"Subject: {msg.get('subject','')}"]
    for part in msg.walk():
        if part.get_content_type() == "text/plain":
            try: chunks.append(part.get_payload(decode=True).decode(part.get_content_charset() or "utf-8", errors="ignore"))
            except Exception: pass
    return ReadResult("\n".join(chunks))


def _msg(data: bytes) -> ReadResult:
    try:
        import extract_msg, tempfile, os
        with tempfile.NamedTemporaryFile(suffix=".msg", delete=False) as f:
            f.write(data); name=f.name
        try:
            m = extract_msg.Message(name)
            return ReadResult(f"Subject: {m.subject or ''}\n{m.body or ''}")
        finally: os.unlink(name)
    except Exception as e:
        return ReadResult("", metadata={"warning": f"MSG reader unavailable: {e}"})


def _image(data: bytes) -> ReadResult:
    img = Image.open(BytesIO(data))
    text = ""
    try:
        import pytesseract
        text = pytesseract.image_to_string(img)
    except Exception:
        pass
    return ReadResult(text, metadata={"image_size": img.size, "needs_ocr": len(text.strip()) < 20})


def read_file(filename: str, data: bytes) -> ReadResult:
    ext = Path(filename).suffix.lower()
    if ext in {".xlsx", ".xlsm", ".xls"}: return _excel(data)
    if ext == ".pdf": return _pdf(data)
    if ext == ".docx": return _docx(data)
    if ext == ".eml": return _eml(data)
    if ext == ".msg": return _msg(data)
    if ext in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"}: return _image(data)
    if ext in {".txt", ".csv"}: return ReadResult(data.decode("utf-8", errors="ignore"))
    raise ValueError(f"Unsupported file type: {ext}")
