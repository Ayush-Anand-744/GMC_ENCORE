from io import BytesIO
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _style(wb):
    for ws in wb.worksheets:
        ws.freeze_panes="A2"
        ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.font=Font(bold=True); cell.fill=PatternFill("solid", fgColor="DCE6F1")
        for col in range(1,ws.max_column+1):
            width=min(max((len(str(ws.cell(r,col).value or "")) for r in range(1,min(ws.max_row,200)+1)),default=10)+2,60)
            ws.column_dimensions[get_column_letter(col)].width=max(12,width)
        for row in ws.iter_rows():
            for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)

def all_results_xlsx(result: dict) -> bytes:
    out=BytesIO()
    with pd.ExcelWriter(out,engine="openpyxl") as w:
        def rows_df(key):
            return pd.DataFrame(result.get(key,[]))
        rows_df("quote_rows").to_excel(w,index=False,sheet_name="Quote Form")
        rows_df("additional_rows").to_excel(w,index=False,sheet_name="Additional Details")
        rows_df("hospital_rows").to_excel(w,index=False,sheet_name="Hospitalization")
        pd.DataFrame(result.get("demography",[])).to_excel(w,index=False,sheet_name="Demography")
        pd.DataFrame(result.get("validation_rows",[])).to_excel(w,index=False,sheet_name="Validation")
        pd.DataFrame(result.get("deviations",[])).to_excel(w,index=False,sheet_name="Deviations")
        _style(w.book)
    return out.getvalue()

def quote_xlsx(result: dict) -> bytes:
    out=BytesIO()
    with pd.ExcelWriter(out,engine="openpyxl") as w:
        pd.DataFrame(result.get("quote_rows",[])).to_excel(w,index=False,sheet_name="Quote Form")
        pd.DataFrame(result.get("hospital_rows",[])).to_excel(w,index=False,sheet_name="Hospitalization")
        pd.DataFrame(result.get("additional_rows",[])).to_excel(w,index=False,sheet_name="Additional Details")
        _style(w.book)
    return out.getvalue()

def demography_xlsx(result: dict) -> bytes:
    out=BytesIO()
    with pd.ExcelWriter(out,engine="openpyxl") as w:
        pd.DataFrame(result.get("demography",[])).to_excel(w,index=False,sheet_name="Demography")
        _style(w.book)
    return out.getvalue()
